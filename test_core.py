"""Offline tests (no network, no LLM). Run:  python -m pytest -q   or   python tests/test_core.py"""
import json
import os
import sys
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd  # noqa: E402

import tracker  # noqa: E402
import verifier  # noqa: E402

TODAY = date(2026, 10, 4)

PAGE = """
Chevening Scholarships are the UK government's international awards.
Applications close on 7 November 2026 at 12:00 GMT.
Results will be announced on 15 March 2027.
The scholarship covers tuition fees, a monthly stipend and one return flight.
Applicants must be citizens of a Chevening-eligible country and hold an undergraduate degree.
Chevening offers one-year master's degrees at UK universities.
"""


def bundle(text=PAGE, ok=True, url="https://www.chevening.org/scholarships/", links=None):
    return {"status": 200 if ok else 404, "ok": ok, "error": "" if ok else "HTTP 404", "final_url": url,
            "text": text, "links": links if links is not None else [("Apply now", "https://www.chevening.org/apply/")],
            "pages": [url]}


REC = {"official_link": "https://www.chevening.org/scholarships/", "provider": "UK FCDO", "scholarship_name": "Chevening Scholarships"}
alive_yes = lambda u: (True, "200")  # noqa: E731
alive_no = lambda u: (False, "404")  # noqa: E731


def run(text=PAGE, level="MS", llm=None, alive=alive_yes, links=None, rec=REC, ok=True, url=None):
    b = bundle(text, ok=ok, links=links, **({"url": url} if url else {}))
    heur = verifier.heuristic_facts(b, rec["official_link"], TODAY)
    return verifier.run_checks(rec, b, llm, heur, level, TODAY, {"chevening.org"}, alive)


# ---------------- dates
def test_find_dates_formats_and_ambiguity():
    t = "Closes 7 November 2026; also Nov 7, 2026; 2026-11-07; 03/04/2027 is ambiguous; 25/12/2026 is clear."
    ds = [d for d, _, _ in verifier.find_dates(t)]
    assert ds.count(date(2026, 11, 7)) == 3 or date(2026, 11, 7) in ds
    assert date(2026, 12, 25) in ds
    assert date(2027, 3, 4) not in ds and date(2027, 4, 3) not in ds


def test_deadline_ignores_results_dates_and_past():
    c = verifier.deadline_candidates(PAGE, TODAY)
    assert [x["date"] for x in c] == ["2026-11-07"]  # results date (2027-03-15) is not a deadline
    old = "Applications closed on 1 March 2025."
    assert verifier.deadline_candidates(old, TODAY) == []


# ---------------- the six checks
def test_all_six_pass():
    r = run()
    assert r["verified"], r["notes"]
    assert r["facts"]["deadline"] == "2026-11-07"
    assert r["facts"]["application_url"] == "https://www.chevening.org/apply/"


def test_aggregator_fails_domain():
    rec = {**REC, "official_link": "https://www.scholars4dev.com/chevening"}
    r = run(rec=rec, url="https://www.scholars4dev.com/chevening", links=[("Apply", "https://www.scholars4dev.com/apply")])
    assert not r["checks"]["official_domain"] and not r["verified"]


def test_dead_page_fails():
    r = run(ok=False)
    assert not r["verified"] and not r["checks"]["official_domain"]


def test_missing_deadline_fails_and_past_flagged():
    r = run(text=PAGE.replace("Applications close on 7 November 2026 at 12:00 GMT.", "Applications are open."))
    assert not r["checks"]["deadline"] and not r["verified"]
    r = run(text=PAGE.replace("7 November 2026", "7 November 2025"))
    assert not r["checks"]["deadline"]


def test_llm_hallucinated_deadline_rejected():
    fake = {"deadline": "2026-12-31", "deadline_evidence": "Applications close on 31 December 2026"}
    r = run(llm=fake)
    assert r["facts"]["deadline"] == "2026-11-07"  # falls back to the date really on the page


def test_llm_fake_funding_quote_rejected():
    fake = {"funding": "Full tuition and 5000 USD monthly", "funding_evidence": "We pay five thousand dollars monthly to every student"}
    r = run(llm=fake)
    assert "5000" not in r["facts"].get("funding", "")


def test_no_funding_no_eligibility_fail():
    t = "Applications close on 7 November 2026. Chevening offers master's degrees."
    r = run(text=t)
    assert not r["checks"]["funding"] and not r["checks"]["eligibility"] and not r["verified"]


def test_wrong_level_fails():
    assert not run(level="PhD")["checks"]["degree_level"]
    assert run(level="MS")["checks"]["degree_level"]


def test_application_url_must_be_https_alive():
    assert not run(alive=alive_no)["checks"]["application_url"]
    r = run(links=[("Apply", "http://www.chevening.org/apply/")])
    assert not r["checks"]["application_url"]
    r = run(links=[])
    assert not r["checks"]["application_url"]


def test_in_text():
    assert verifier.in_text("monthly stipend and one return flight", PAGE)
    assert not verifier.in_text("a completely invented sentence about money", PAGE)


# ---------------- tracker
def _df(days):
    rows = []
    for i, d in enumerate(days):
        dl = (date.today() + timedelta(days=d)).isoformat() if d is not None else None
        rows.append(dict(id=i + 1, country="UK", scholarship_name=f"S{i}", provider="p", level="MS", deadline=dl,
                         status="Remaining", notes="", fit_score=50, verified=True, official_link=f"https://x.org/{i}",
                         application_url="", verification="✓✓✓✓✓✓", checks_json="{}", notes_json="{}"))
    return tracker.recompute(pd.DataFrame(rows))


def test_urgency_levels():
    df = _df([3, 20, 100, -5, None])
    assert list(df["urgency"]) == ["Critical", "Soon", "OK", "Closed", "Verify"]
    assert list(tracker.critical_list(df)["scholarship_name"]) == ["S0"]
    assert tracker.status_counts(df)["critical"] == 1


def test_applied_not_critical():
    df = _df([3])
    df.loc[0, "status"] = "Applied"
    assert tracker.critical_list(df).empty


def test_no_fuzzy_deadline_parsing():
    assert tracker.parse_deadline("sometime in spring") is None
    assert tracker.parse_deadline("7 November 2026") == date(2026, 11, 7)


def test_ics():
    ics = tracker.to_ics(_df([3, None]))
    assert ics.count("BEGIN:VEVENT") == 1 and "TRIGGER:-P14D" in ics and "TRIGGER:-P3D" in ics


def test_excel_export():
    import export
    from io import BytesIO
    from openpyxl import load_workbook

    df = _df([3, 40])
    b = export.to_excel_bytes(df, None)
    wb = load_workbook(BytesIO(b))
    assert wb.sheetnames == ["Scholarships", "Gap Report", "Verification"] or set(wb.sheetnames) >= {"Scholarships", "Verification"}


if __name__ == "__main__":
    fails = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print("PASS", name)
            except Exception as exc:  # noqa: BLE001
                fails += 1
                print("FAIL", name, type(exc).__name__, exc)
    sys.exit(1 if fails else 0)
